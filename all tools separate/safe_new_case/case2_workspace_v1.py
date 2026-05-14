from pathlib import Path
import json
from io import BytesIO
from datetime import datetime

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from docx import Document

from case2_engine_v1 import (
    WORKSPACE_CONTEXT_PATH,
    build_case2_workspace_view,
)

# ── PAGE CONFIG ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Case 2 · Strategic Sourcing Workspace",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── THEME / GLOBAL CSS ───────────────────────────────────────────────────────
st.markdown(
    """
    <style>
    /* ---- Fonts ---- */
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap');

    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }

    /* ---- Page background ---- */
    .stApp {
        background-color: #F4F6F9;
    }

    /* ---- Sidebar ---- */
    section[data-testid="stSidebar"] {
        background-color: #1A2235;
        border-right: 1px solid #2C3A52;
    }
    section[data-testid="stSidebar"] * {
        color: #C8D3E6 !important;
    }
    section[data-testid="stSidebar"] .stSelectbox label,
    section[data-testid="stSidebar"] .stMultiSelect label {
        font-size: 11px;
        font-weight: 600;
        letter-spacing: 0.07em;
        text-transform: uppercase;
        color: #7A8FAF !important;
        margin-bottom: 2px;
    }
    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {
        color: #E8EDF5 !important;
        font-size: 13px !important;
        font-weight: 600 !important;
        letter-spacing: 0.05em;
        text-transform: uppercase;
        margin-bottom: 12px;
    }
    /* Sidebar select + multiselect inputs */
    section[data-testid="stSidebar"] .stSelectbox > div > div,
    section[data-testid="stSidebar"] .stMultiSelect > div > div {
        background-color: #232F45 !important;
        border: 1px solid #2C3A52 !important;
        border-radius: 6px !important;
        color: #C8D3E6 !important;
    }
    /* Multiselect tags */
    section[data-testid="stSidebar"] span[data-baseweb="tag"] {
        background-color: #1E4A8A !important;
        border-radius: 4px !important;
    }
    section[data-testid="stSidebar"] span[data-baseweb="tag"] span {
        color: #BFCFE8 !important;
        font-size: 11px !important;
    }
    /* Divider in sidebar */
    section[data-testid="stSidebar"] hr {
        border-color: #2C3A52;
        margin: 16px 0;
    }

    /* ---- KPI metric cards ---- */
    div[data-testid="metric-container"] {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 10px;
        padding: 18px 20px 14px 20px;
        box-shadow: 0 1px 4px rgba(0,0,0,0.05);
    }
    div[data-testid="metric-container"] label {
        font-size: 11px !important;
        font-weight: 600 !important;
        letter-spacing: 0.07em !important;
        text-transform: uppercase !important;
        color: #64748B !important;
    }
    div[data-testid="metric-container"] div[data-testid="stMetricValue"] {
        font-family: 'IBM Plex Mono', monospace !important;
        font-size: 24px !important;
        font-weight: 500 !important;
        color: #0F172A !important;
    }
    div[data-testid="metric-container"] div[data-testid="stMetricDelta"] {
        font-size: 12px !important;
    }

    /* ---- Chart cards ---- */
    div[data-testid="stPlotlyChart"] {
        background: #FFFFFF;
        border: 1px solid #E2E8F0;
        border-radius: 10px;
        padding: 8px;
        box-shadow: 0 1px 4px rgba(0,0,0,0.05);
    }

    /* ---- Insight banner ---- */
    div[data-testid="stAlert"] {
        border-radius: 8px !important;
        border-left: 4px solid #2563EB !important;
        background: #EFF6FF !important;
        font-size: 13px !important;
        color: #1E3A5F !important;
        padding: 12px 16px !important;
    }

    /* ---- Success banner (agentic focus) ---- */
    div[data-testid="stNotification"] {
        border-radius: 8px !important;
    }

    /* ---- Caption / locked labels ---- */
    .locked-label {
        font-size: 11px;
        color: #94A3B8;
        font-family: 'IBM Plex Mono', monospace;
        margin-top: -6px;
        margin-bottom: 10px;
    }

    /* ---- Section divider ---- */
    .section-rule {
        border: none;
        border-top: 1px solid #E2E8F0;
        margin: 24px 0 16px 0;
    }

    /* ---- Page title area ---- */
    .ws-header {
        padding: 4px 0 20px 0;
        border-bottom: 1px solid #E2E8F0;
        margin-bottom: 24px;
    }
    .ws-title {
        font-size: 22px;
        font-weight: 600;
        color: #0F172A;
        letter-spacing: -0.3px;
    }
    .ws-subtitle {
        font-size: 13px;
        color: #64748B;
        margin-top: 3px;
    }
    .ws-tag {
        display: inline-block;
        background: #DBEAFE;
        color: #1D4ED8;
        font-size: 11px;
        font-weight: 600;
        letter-spacing: 0.05em;
        text-transform: uppercase;
        border-radius: 4px;
        padding: 2px 8px;
        margin-right: 8px;
        vertical-align: middle;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ── PLOTLY SHARED THEME ──────────────────────────────────────────────────────
CHART_FONT = dict(family="Inter, sans-serif", size=12, color="#374151")
CHART_COLORS = {
    "active":     "#1D4ED8",
    "future":     "#93C5FD",
    "cost_line":  "#EF4444",
    "proposed":   "#1D4ED8",
    "market":     "#64748B",
    "competitor": "#EF4444",
    "tiers": {
        "Standard": "#93C5FD",
        "Pro":      "#1D4ED8",
        "Premium":  "#1E3A8A",
    },
    "donut": ["#1D4ED8","#3B82F6","#93C5FD","#BFDBFE","#60A5FA","#2563EB","#DBEAFE"],
}

LAYOUT_BASE = dict(
    font=CHART_FONT,
    paper_bgcolor="rgba(0,0,0,0)",
    plot_bgcolor="rgba(0,0,0,0)",
    margin=dict(l=48, r=24, t=72, b=40),
    legend=dict(
        orientation="h",
        yanchor="bottom",
        y=1.18,
        xanchor="left",
        x=0,
        font=dict(size=11),
        bgcolor="rgba(0,0,0,0)",
        borderwidth=0,
    ),
    hoverlabel=dict(
        bgcolor="#1A2235",
        font_color="#E8EDF5",
        font_size=12,
        bordercolor="#2C3A52",
    ),
)

AXIS_STYLE = dict(
    showgrid=True,
    gridcolor="#F1F5F9",
    gridwidth=1,
    zeroline=False,
    linecolor="#E2E8F0",
    tickfont=dict(size=11, color="#6B7280"),
    title_font=dict(size=12, color="#374151"),
)


def styled_layout(title: str, **kwargs) -> dict:
    layout = {**LAYOUT_BASE, "title": dict(text=title, font=dict(size=14, weight=600, color="#0F172A"), x=0, xanchor="left", y=0.99, yanchor="top"), **kwargs}
    return layout


# ── DATA HELPERS ─────────────────────────────────────────────────────────────

def load_workspace_context() -> dict:
    if WORKSPACE_CONTEXT_PATH.exists():
        with open(WORKSPACE_CONTEXT_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def build_sourcing_insight(filtered_df: pd.DataFrame, company_df: pd.DataFrame) -> str:
    if filtered_df.empty or company_df.empty:
        return "Insufficient data to generate a tier comparison for this selection."
    filtered_premium = (filtered_df["tier_group"] == "Premium").mean()
    company_premium  = (company_df["tier_group"]  == "Premium").mean()
    delta            = filtered_premium - company_premium
    filtered_pct     = filtered_premium * 100
    company_pct      = company_premium  * 100
    total_users      = int(filtered_df["employee_id"].nunique())

    if delta > 0.05:
        return (
            f"Premium tier concentration is elevated for this selection. "
            f"{filtered_pct:.1f}% of users in scope are mapped to Premium tiers versus "
            f"{company_pct:.1f}% at company baseline (a +{delta*100:.1f} pp gap across {total_users:,} users). "
            f"Consider whether all targeted personas require Premium entitlements, "
            f"or whether a portion of demand can be shifted to lower tiers to reduce spend concentration."
        )
    if delta < -0.05:
        return (
            f"Premium tier concentration is below the company baseline. "
            f"{filtered_pct:.1f}% of users in scope are mapped to Premium tiers versus "
            f"{company_pct:.1f}% baseline (a \u2212{abs(delta)*100:.1f} pp gap across {total_users:,} users). "
            f"This selection is cost-efficient, but confirm that critical roles still receive "
            f"advanced capabilities where operationally required."
        )
    return (
        f"Tier mix is broadly aligned with company norms. "
        f"Premium coverage is {filtered_pct:.1f}% in scope versus {company_pct:.1f}% baseline "
        f"({delta*100:+.1f} pp across {total_users:,} users). "
        f"Current allocation appears balanced across demand and spend."
    )


def build_negotiation_delta_insight(filtered_df: pd.DataFrame) -> str:
    if filtered_df.empty:
        return "Insufficient data to assess negotiation delta for this selection."

    spend_frame = (
        filtered_df.groupby("target_sku_id", dropna=False)
        .agg(
            users=("employee_id", "nunique"),
            proposed_unit_price=("resolved_unit_price", "first"),
            market_avg_price=("market_avg_price", "first"),
            discount_percent=("discount_percent", "first"),
        )
        .reset_index()
    )
    if spend_frame.empty:
        return "Insufficient pricing data to assess negotiation delta for this selection."

    proposed_spend = float((spend_frame["users"] * spend_frame["proposed_unit_price"] * 12).sum())
    market_spend = float((spend_frame["users"] * spend_frame["market_avg_price"] * 12).sum())
    delta = proposed_spend - market_spend
    total_users = int(filtered_df["employee_id"].nunique())
    gap_pct = (delta / market_spend * 100) if market_spend else 0.0
    weighted_discount_pct = (
        float((spend_frame["users"] * spend_frame["discount_percent"]).sum()) / float(spend_frame["users"].sum())
        if not spend_frame.empty and float(spend_frame["users"].sum()) > 0
        else 0.0
    )
    tier_context = (
        f"Current scoped pricing already includes volume-band discounts (weighted avg: {weighted_discount_pct:.1f}% off list). "
    )

    if delta > 0:
        return (
            f"Negotiation Delta Insight · Proposed annual spend is above market for this selection. "
            + tier_context
            + f"In-scope projected spend is ${proposed_spend:,.0f} versus ${market_spend:,.0f} at market baseline "
            + f"(a +${delta:,.0f} gap, +{gap_pct:.1f}% across {total_users:,} users). "
            + "Consider repricing high-volume SKUs toward market benchmarks before final commitment."
        )
    if delta < 0:
        return (
            f"Negotiation Delta Insight · Current pricing is favorable versus market for this selection. "
            + tier_context
            + f"In-scope projected spend is ${proposed_spend:,.0f} versus ${market_spend:,.0f} at market baseline "
            + f"(a -${abs(delta):,.0f} gap, {gap_pct:.1f}% across {total_users:,} users). "
            + "Capture this advantage in term commitments while preserving flexibility for volume changes."
        )
    return (
        f"Negotiation Delta Insight · Proposed pricing is aligned with market for this selection. "
        + tier_context
        + f"In-scope projected spend matches the market baseline at ${proposed_spend:,.0f} "
        + f"(~${delta:,.0f} variance across {total_users:,} users)."
    )


def build_tier_alias_map(df: pd.DataFrame) -> dict[str, str]:
    """
    Normalize raw license tiers to chart-aligned tier groups for uniform filtering.
    Example mappings observed from data: Full->Pro, Editor/Viewer->Standard, Enterprise->Premium.
    """
    if df.empty:
        return {}
    mapping_df = (
        df.dropna(subset=["target_license_tier", "tier_group"])
        .groupby("target_license_tier")["tier_group"]
        .agg(lambda s: s.mode().iloc[0] if not s.mode().empty else s.iloc[0])
        .reset_index()
    )
    return dict(zip(mapping_df["target_license_tier"], mapping_df["tier_group"]))


def build_dummy_gemma_summary(filtered_df: pd.DataFrame, selected_vendor: str | None) -> str:
    """
    Dummy Gemma-style summary text.
    Replace this with a real LLM call in the next iteration.
    """
    if filtered_df.empty:
        return (
            "No records are currently in scope for this request. "
            "Please adjust filters and regenerate the report."
        )

    total_users = int(filtered_df["employee_id"].nunique())
    future_hires = int((filtered_df["employee_status"] == "Future_Hire").sum())
    annual_spend = float(filtered_df["projected_annual_cost"].sum())
    top_dept = (
        filtered_df.groupby("pillar_dept").size().sort_values(ascending=False).index.tolist()
    )
    top_sku = (
        filtered_df.groupby("target_sku_name").size().sort_values(ascending=False).index.tolist()
    )
    tier_mix = (
        filtered_df.groupby("tier_group").size().sort_values(ascending=False)
    )
    tier_line = ", ".join([f"{k}: {v}" for k, v in tier_mix.items()])

    return (
        f"Gemma draft summary for new procurement request ({selected_vendor or 'selected vendor'}).\n"
        f"The current demand scope includes {total_users:,} users with {future_hires:,} future hires, "
        f"driving an estimated annual spend of ${annual_spend:,.0f}. "
        f"Demand concentration is led by department: {top_dept[0] if top_dept else 'N/A'}, "
        f"and primary license candidate: {top_sku[0] if top_sku else 'N/A'}. "
        f"Observed tier distribution in scope: {tier_line if tier_line else 'N/A'}. "
        "Recommendation: proceed with phased procurement, validate Premium entitlement fit for high-cost cohorts, "
        "and use negotiated pricing benchmarks before final purchase commitment."
    )


def build_report_docx_bytes(
    filtered_df: pd.DataFrame,
    company_df: pd.DataFrame,
    selected_vendor: str | None,
    selected_dept: str,
    selected_licenses: list[str],
    selected_tiers: list[str],
) -> bytes:
    doc = Document()
    doc.add_heading("Case 2 Procurement Summary Report", level=0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph("Prompt context: I want to buy new licenses from vendor_3.")

    doc.add_heading("Scope", level=1)
    doc.add_paragraph(f"Vendor: {selected_vendor or 'N/A'}")
    doc.add_paragraph(f"Department: {selected_dept}")
    doc.add_paragraph(f"Licenses: {', '.join(selected_licenses) if selected_licenses else 'All'}")
    doc.add_paragraph(f"License Tiers: {', '.join(selected_tiers) if selected_tiers else 'All'}")

    doc.add_heading("Gemma Summary (Dummy)", level=1)
    doc.add_paragraph(build_dummy_gemma_summary(filtered_df, selected_vendor))
    doc.add_heading("Sourcing Insights", level=1)
    doc.add_paragraph("Sourcing Insight · " + build_sourcing_insight(filtered_df, company_df))
    doc.add_paragraph(build_negotiation_delta_insight(filtered_df))

    if not filtered_df.empty:
        doc.add_heading("Top License Demand Snapshot", level=1)
        snapshot = (
            filtered_df.groupby(["target_sku_name", "tier_group"], dropna=False)
            .agg(
                users=("employee_id", "nunique"),
                annual_spend=("projected_annual_cost", "sum"),
            )
            .reset_index()
            .sort_values("users", ascending=False)
            .head(8)
        )
        table = doc.add_table(rows=1, cols=4)
        hdr = table.rows[0].cells
        hdr[0].text = "License"
        hdr[1].text = "Tier Group"
        hdr[2].text = "Users"
        hdr[3].text = "Annual Spend"
        for _, row in snapshot.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["target_sku_name"])
            cells[1].text = str(row["tier_group"])
            cells[2].text = f"{int(row['users']):,}"
            cells[3].text = f"${float(row['annual_spend']):,.0f}"

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


# ── LOAD DATA ────────────────────────────────────────────────────────────────
workspace_df = build_case2_workspace_view()
context      = load_workspace_context()
defaults     = context.get("workspace_defaults", {}) if isinstance(context, dict) else {}

focus_sku_id             = defaults.get("focus_sku_id")
focus_sku_name           = defaults.get("focus_sku_name")
lock_to_requested_license = bool(defaults.get("lock_to_requested_license"))

# ── PAGE HEADER ──────────────────────────────────────────────────────────────
st.title("Agentic Demand Discovery Workspace")
st.caption("Case 2 · Demand discovery and pricing analysis for new license procurement")
if focus_sku_name and defaults.get("vendor_name"):
    st.caption(f"Agentic Focus · {focus_sku_name} · {defaults.get('vendor_name')}")
st.divider()

# ── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### Filters")
    st.markdown("<hr/>", unsafe_allow_html=True)

    # Vendor
    vendor_options = sorted(workspace_df["target_vendor"].dropna().unique().tolist())
    default_vendor = defaults.get("vendor_name") if defaults.get("vendor_name") in vendor_options else None
    if default_vendor:
        selected_vendor = default_vendor
        st.markdown(f"**Vendor**")
        st.markdown(f"<div class='locked-label'>{selected_vendor} &nbsp;·&nbsp; locked from request</div>", unsafe_allow_html=True)
    else:
        selected_vendor = st.selectbox("Vendor", vendor_options, index=0) if vendor_options else None

    vendor_scoped_df = workspace_df[workspace_df["target_vendor"] == selected_vendor].copy() if selected_vendor else workspace_df.copy()

    # Department
    dept_options   = ["All"] + sorted(vendor_scoped_df["pillar_dept"].dropna().unique().tolist())
    default_dept   = defaults.get("pillar_dept") if defaults.get("pillar_dept") in dept_options else "All"
    selected_dept  = st.selectbox("Department", dept_options, index=dept_options.index(default_dept))

    st.markdown("<hr/>", unsafe_allow_html=True)

    # Build filtered base for downstream selects
    filtered = workspace_df.copy()
    if selected_vendor:
        filtered = filtered[filtered["target_vendor"] == selected_vendor]
    if focus_sku_id:
        filtered = filtered[filtered["target_sku_id"] == focus_sku_id]
    elif focus_sku_name:
        filtered = filtered[filtered["target_sku_name"].str.lower().eq(str(focus_sku_name).lower())]
    if selected_dept != "All":
        filtered = filtered[filtered["pillar_dept"] == selected_dept]

    # License name
    license_options = sorted(filtered["target_sku_name"].dropna().unique().tolist())
    if lock_to_requested_license and focus_sku_name:
        selected_licenses = [focus_sku_name] if focus_sku_name in license_options else []
        st.markdown(f"**License**")
        st.markdown(f"<div class='locked-label'>{focus_sku_name} &nbsp;·&nbsp; locked from request</div>", unsafe_allow_html=True)
    else:
        selected_licenses = st.multiselect(
            "License",
            options=license_options,
            default=license_options,
            key=f"lic_{selected_vendor}_{selected_dept}_{focus_sku_id}_{focus_sku_name}",
        )
    if selected_licenses:
        filtered = filtered[filtered["target_sku_name"].isin(selected_licenses)]

    # License tier (normalized aliases aligned with charts: Standard / Pro / Premium)
    filtered = filtered.copy()
    tier_alias_map = build_tier_alias_map(filtered)
    filtered["normalized_tier"] = (
        filtered["target_license_tier"]
        .map(tier_alias_map)
        .fillna(filtered["tier_group"])
        .fillna(filtered["target_license_tier"])
    )
    tier_counts          = filtered.groupby("normalized_tier").size().sort_values(ascending=False)
    tier_options         = tier_counts.index.tolist()
    selected_tiers       = st.multiselect(
        "License Tier",
        options=tier_options,
        default=tier_options,
        key=f"tier_{selected_vendor}_{selected_dept}_{focus_sku_id}_{focus_sku_name}",
        help="Normalized tier aliases for consistent filtering and chart interpretation.",
    )
    if selected_tiers:
        filtered = filtered[filtered["normalized_tier"].isin(selected_tiers)]

    # Job role
    role_counts    = filtered.groupby("job_role").size().sort_values(ascending=False)
    role_options   = role_counts.index.tolist()
    default_roles  = role_options.copy()
    if not default_roles:
        default_roles = [r for r in defaults.get("top_roles", []) if r in role_options]

    selected_roles = st.multiselect(
        "Job Role",
        options=role_options,
        default=default_roles,
        key=f"role_{selected_vendor}_{selected_dept}",
    )
    if selected_roles:
        filtered = filtered[filtered["job_role"].isin(selected_roles)]

    st.markdown("<hr/>", unsafe_allow_html=True)
    st.caption(f"{len(filtered):,} employee records in scope")

# ── KPIs ─────────────────────────────────────────────────────────────────────
total_demand  = int(filtered["employee_id"].count())
future_demand = int((filtered["employee_status"] == "Future_Hire").sum())

spend_frame = (
    filtered.groupby("target_sku_id", dropna=False)
    .agg(
        users=("employee_id", "nunique"),
        resolved_unit_price=("resolved_unit_price", "first"),
        market_avg_price=("market_avg_price", "first"),
    )
    .reset_index()
)
est_annual_spend  = float((spend_frame["users"] * spend_frame["resolved_unit_price"] * 12).sum())
negotiation_delta = float(((spend_frame["resolved_unit_price"] - spend_frame["market_avg_price"]) * spend_frame["users"] * 12).sum())

k1, k2, k3, k4 = st.columns(4)
k1.metric("Total Project Demand",  f"{total_demand:,}",  delta=f"{future_demand:,} future hires")
k2.metric("Active Employees",      f"{total_demand - future_demand:,}")
k3.metric("Est. Annual Spend",     f"${est_annual_spend:,.0f}")
k4.metric("Negotiation Delta",     f"${negotiation_delta:,.0f}", delta=f"{'above' if negotiation_delta > 0 else 'below'} market avg", delta_color="inverse")

# ── SOURCING INSIGHT ─────────────────────────────────────────────────────────
st.info("**Sourcing Insight** · " + build_sourcing_insight(filtered, workspace_df))
st.info(build_negotiation_delta_insight(filtered))

st.markdown("<div class='section-rule'></div>", unsafe_allow_html=True)

# ── CHARTS ───────────────────────────────────────────────────────────────────
col_left, col_right = st.columns(2, gap="medium")

# Chart 1 — Aggregated Demand & Cost
with col_left:
    active_count = int((filtered["employee_status"] == "Active").sum())
    future_count = int((filtered["employee_status"] == "Future_Hire").sum())

    fig1 = go.Figure()
    fig1.add_bar(
        x=["Total Project"],
        y=[active_count],
        name="Active",
        marker_color=CHART_COLORS["active"],
        hovertemplate="<b>Active Employees</b><br>Count: %{y:,}<extra></extra>",
    )
    fig1.add_bar(
        x=["Total Project"],
        y=[future_count],
        name="Future Hire",
        marker_color=CHART_COLORS["future"],
        hovertemplate="<b>Future Hires</b><br>Count: %{y:,}<extra></extra>",
    )
    fig1.add_trace(go.Scatter(
        x=["Total Project"],
        y=[est_annual_spend],
        yaxis="y2",
        mode="markers",
        marker=dict(symbol="diamond", size=12, color=CHART_COLORS["cost_line"], line=dict(width=2, color="#FFFFFF")),
        name="Annual Cost",
        hovertemplate="<b>Annual Cost</b><br>$%{y:,.0f}<extra></extra>",
    ))
    fig1.update_layout(
        **styled_layout("Aggregated Demand & Cost"),
        barmode="stack",
        yaxis=dict(title="Headcount", **AXIS_STYLE),
        yaxis2=dict(title="Annual Cost ($)", overlaying="y", side="right", showgrid=False, tickfont=dict(size=11, color="#6B7280"), title_font=dict(size=12, color="#374151"), zeroline=False),
        showlegend=True,
    )
    st.plotly_chart(fig1, use_container_width=True)

# Chart 2 — Departmental Persona Mix
with col_right:
    donut_df = (
        filtered.groupby("pillar_dept", dropna=False)
        .agg(count=("employee_id", "count"), annual_spend=("projected_annual_cost", "sum"))
        .reset_index()
        .sort_values("count", ascending=False)
    )

    fig2 = go.Figure(go.Pie(
        labels=donut_df["pillar_dept"],
        values=donut_df["count"],
        customdata=donut_df[["annual_spend"]],
        hole=0.58,
        marker=dict(colors=CHART_COLORS["donut"], line=dict(color="#FFFFFF", width=2)),
        hovertemplate=(
            "<b>%{label}</b><br>"
            "Headcount: %{value:,}<br>"
            "Annual Spend: $%{customdata[0]:,.0f}<br>"
            "Share: %{percent}<extra></extra>"
        ),
        textposition="outside",
        textfont=dict(size=11),
    ))
    fig2_layout = styled_layout("Departmental Persona Mix")
    fig2_layout["legend"] = dict(orientation="v", yanchor="middle", y=0.5, xanchor="left", x=1.02, font=dict(size=11))
    fig2_layout["showlegend"] = True
    fig2.update_layout(**fig2_layout)
    st.plotly_chart(fig2, use_container_width=True)

st.markdown("<div class='section-rule'></div>", unsafe_allow_html=True)

col_left2, col_right2 = st.columns(2, gap="medium")

# Chart 3 — Tier Optimisation
with col_left2:
    tier_df = (
        filtered.groupby(["job_role", "tier_group"], dropna=False)
        .agg(count=("employee_id", "count"), resolved_unit_price=("resolved_unit_price", "first"))
        .reset_index()
    )
    tier_df["stack_cost"] = tier_df["count"] * tier_df["resolved_unit_price"] * 12

    # Top 4 + Others logic
    role_totals  = tier_df.groupby("job_role")["count"].sum().sort_values(ascending=False)
    top4         = role_totals.head(4).index.tolist()
    tier_df["role_label"] = tier_df["job_role"].apply(lambda r: r if r in top4 else "Others")
    tier_plot    = tier_df.groupby(["role_label", "tier_group"], dropna=False).agg(count=("count","sum"), stack_cost=("stack_cost","sum"), resolved_unit_price=("resolved_unit_price","first")).reset_index()
    # Preserve sort order: top4 first, Others last
    order        = top4 + (["Others"] if "Others" in tier_plot["role_label"].values else [])
    tier_plot["_ord"] = tier_plot["role_label"].apply(lambda r: order.index(r) if r in order else 99)
    tier_plot    = tier_plot.sort_values("_ord")

    fig3 = go.Figure()
    for tier_key, color in CHART_COLORS["tiers"].items():
        sub = tier_plot[tier_plot["tier_group"] == tier_key]
        if sub.empty:
            continue
        fig3.add_bar(
            x=sub["role_label"],
            y=sub["count"],
            name=tier_key,
            marker_color=color,
            customdata=sub[["stack_cost", "resolved_unit_price"]],
            hovertemplate=(
                f"<b>{{%{{x}}}}</b> · {tier_key}<br>"
                "Headcount: %{y:,}<br>"
                "Stack Cost: $%{customdata[0]:,.0f}/yr<br>"
                "Unit Price: $%{customdata[1]:,.2f}/mo<extra></extra>"
            ),
        )
    fig3.update_layout(
        **styled_layout("Tier Optimisation · Top 4 Roles + Others"),
        barmode="stack",
        xaxis=dict(title="Job Role", **AXIS_STYLE),
        yaxis=dict(title="Headcount", **AXIS_STYLE),
    )
    st.plotly_chart(fig3, use_container_width=True)

# Chart 4 — Market Sourcing Analysis
with col_right2:
    market_df = (
        filtered.groupby(["target_sku_id", "target_sku_name"], dropna=False)
        .agg(
            proposed_price=("resolved_unit_price", "first"),
            market_avg=("market_avg_price", "first"),
            competitor_price=("best_competitor_price", "first"),
        )
        .reset_index()
    )
    if lock_to_requested_license and focus_sku_name:
        market_df = market_df[market_df["target_sku_name"].str.lower().eq(str(focus_sku_name).lower())]

    fig4 = go.Figure()
    fig4.add_bar(
        x=market_df["target_sku_name"],
        y=market_df["proposed_price"],
        name="Proposed Price",
        marker_color=CHART_COLORS["proposed"],
        hovertemplate="<b>%{x}</b><br>Proposed: $%{y:,.2f}<extra></extra>",
    )
    fig4.add_bar(
        x=market_df["target_sku_name"],
        y=market_df["market_avg"],
        name="Market Average",
        marker_color=CHART_COLORS["market"],
        hovertemplate="<b>%{x}</b><br>Market Avg: $%{y:,.2f}<extra></extra>",
    )
    fig4.add_bar(
        x=market_df["target_sku_name"],
        y=market_df["competitor_price"],
        name="Best Competitor",
        marker_color=CHART_COLORS["competitor"],
        hovertemplate="<b>%{x}</b><br>Best Competitor: $%{y:,.2f}<extra></extra>",
    )
    fig4.update_layout(
        **styled_layout("Market Sourcing Analysis"),
        barmode="group",
        xaxis=dict(title="Vendor SKU", **AXIS_STYLE, tickangle=-30),
        yaxis=dict(title="Unit Price ($/mo)", **AXIS_STYLE),
    )
    st.plotly_chart(fig4, use_container_width=True)

# ── REPORT GENERATION / DOWNLOAD ─────────────────────────────────────────────
st.markdown("<div class='section-rule'></div>", unsafe_allow_html=True)
st.subheader("Procurement Report")
st.caption("Generate and download a dummy Gemma-style summary report for the current dashboard scope.")

if st.button("Generate Report (Dummy Gemma Summary)"):
    report_bytes = build_report_docx_bytes(
        filtered_df=filtered,
        company_df=workspace_df,
        selected_vendor=selected_vendor,
        selected_dept=selected_dept,
        selected_licenses=selected_licenses,
        selected_tiers=selected_tiers,
    )
    st.session_state["case2_report_docx_bytes"] = report_bytes
    st.success("Report generated. Use the download button below.")

if "case2_report_docx_bytes" in st.session_state:
    report_name = f"case2_procurement_report_{(selected_vendor or 'vendor').lower()}_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
    st.download_button(
        "Download Procurement Report (.docx)",
        data=st.session_state["case2_report_docx_bytes"],
        file_name=report_name,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

# ── FOOTER ───────────────────────────────────────────────────────────────────
st.markdown("<div class='section-rule'></div>", unsafe_allow_html=True)
st.caption(" Demand Discovery Workspace · Demand Planning Engine · Confidential")